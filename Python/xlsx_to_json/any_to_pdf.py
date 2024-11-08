# pip install fpdf pandas openpyxl python-docx

import os
import shutil
import pandas as pd
import json

from pathlib import Path
from fpdf import FPDF
from docx import Document

class PDF(FPDF):
    def footer(self):
        self.set_y(-15)
        self.set_font('Arial', 'I', 8)
        self.cell(0, 10, f'Page {self.page_no()}', 0, 0, 'C')

def clear_out(out_folder):
    if os.path.exists(out_folder):
        shutil.rmtree(out_folder)
    os.makedirs(out_folder)

def convert_txt_to_pdf(input_file, output_file):
    pdf = PDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Arial", size=10)
    encodings = ['utf-8', 'latin1', 'iso-8859-1']
    for encoding in encodings:
        try:
            with open(input_file, 'r', encoding=encoding) as file:
                for line in file:
                    pdf.multi_cell(0, 10, line.encode('latin1', 'replace').decode('latin1'))
            break
        except UnicodeDecodeError:
            continue
    pdf.output(output_file)

def convert_xlsx_to_pdf(input_file, output_file):
    try:
        df = pd.read_excel(input_file, engine='openpyxl')
        pdf = PDF()
        pdf.add_page()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.set_font("Arial", size=8)

        col_width = pdf.w / (len(df.columns) + 1)
        row_height = pdf.font_size * 1.5

        for col in df.columns:
            pdf.cell(col_width, row_height, str(col), border=1)
        pdf.ln(row_height)

        for i in range(len(df)):
            row = df.iloc[i]
            for item in row:
                pdf.cell(col_width, row_height, str(item), border=1)
            pdf.ln(row_height)

        pdf.output(output_file)
    except Exception as e:
        print(f"Erro ao converter {input_file}: {e}")

def convert_csv_to_pdf(input_file, output_file):
    df = pd.read_csv(input_file)
    pdf = PDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Arial", size=8)

    col_width = pdf.w / (len(df.columns) + 1)
    row_height = pdf.font_size * 1.5

    for col in df.columns:
        pdf.cell(col_width, row_height, str(col), border=1)
    pdf.ln(row_height)

    for i in range(len(df)):
        row = df.iloc[i]
        for item in row:
            pdf.cell(col_width, row_height, str(item), border=1)
        pdf.ln(row_height)

    pdf.output(output_file)

def convert_json_to_pdf(input_file, output_file):
    with open(input_file, 'r', encoding='utf-8') as file:
        data = json.load(file)
    pdf = PDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Arial", size=10)
    pdf.multi_cell(0, 10, json.dumps(data, indent=4))
    pdf.output(output_file)

def convert_docx_to_pdf(input_file, output_file):
    doc = Document(input_file)
    pdf = PDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Arial", size=10)
    for para in doc.paragraphs:
        pdf.multi_cell(0, 10, para.text)
    pdf.output(output_file)

def main():
    input_folder = 'in'
    out_folder = 'out'
    clear_out(out_folder)

    for root, _, files in os.walk(input_folder):
        for file in files:
            input_file = os.path.join(root, file)
            output_file = os.path.join(out_folder, Path(file).stem + '.pdf')
            if file.endswith('.txt'):
                convert_txt_to_pdf(input_file, output_file)
            elif file.endswith('.xlsx'):
                convert_xlsx_to_pdf(input_file, output_file)
            elif file.endswith('.csv'):
                convert_csv_to_pdf(input_file, output_file)
            elif file.endswith('.json'):
                convert_json_to_pdf(input_file, output_file)
            elif file.endswith('.docx'):
                convert_docx_to_pdf(input_file, output_file)

if __name__ == "__main__":
    main()